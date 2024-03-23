from docx import Document
import PyPDF2
from openpyxl import load_workbook
import os


def extraer_docx_metadata(archivo_ruta):
    doc = Document(archivo_ruta)
    metadata = {}
    core_props = doc.core_properties
    metadata['Titulo'] = core_props.title
    metadata['Autor'] = core_props.author
    metadata['Sujeto'] = core_props.subject
    metadata['Palabra_clave'] = core_props.keywords
    metadata['Creado'] = core_props.created
    metadata['Modificado'] = core_props.modified
    return metadata


def extract_xlsx_metadata(archivo_ruta):
    wb = load_workbook(archivo_ruta)
    metadata = {}
    props = wb.properties
    metadata['Titulo'] = props.title
    metadata['Autor'] = props.creator
    metadata['Sujeto'] = props.subject
    metadata['Creado'] = props.created
    metadata['Modificado'] = props.modified
    return metadata



def extract_pdf_metadata(archivo_ruta):
    with open(archivo_ruta, 'rb') as archivo:
        pdf_lectura = PyPDF2.PdfReader(archivo)
        metadata = pdf_lectura.metadata
    return metadata



# Ruta de la carpeta
folder_ruta = '/home/jcsr/Documents/archivos_prueba/'

# Iterar sobre los archivos en la carpeta
for archvio_nombre in os.listdir(folder_ruta):
    archivo_ruta = os.path.join(folder_ruta, archvio_nombre)
    if os.path.isfile(archivo_ruta):
        _,archivo_ext=os.path.splitext(archvio_nombre)
        if archivo_ext==".docx":
            print("============Metadata de {}===========================".format(archvio_nombre))
            metadata = extraer_docx_metadata(archivo_ruta)
            print(metadata)
        elif archivo_ext==".xlsx":
            print("============Metadata de {}===========================".format(archvio_nombre))
            metadata = extract_xlsx_metadata(archivo_ruta)
            print(metadata)
        elif archivo_ext==".pdf":
            print("============Metadata de {}===========================".format(archvio_nombre))
            metadata = extract_pdf_metadata(archivo_ruta)
            print(metadata)
        else:
            pass
    elif os.path.isdir(archivo_ruta):
        pass

